"""Generate Phase 12 dry-run fixture documents."""

from __future__ import annotations

import json
from pathlib import Path

import pymupdf


def _write_pdf(path: Path, pages: list[str]) -> None:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text, fontsize=11)
    doc.save(path)
    doc.close()


def _write_table_pdf(path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Product Specifications", fontsize=14)
    page.insert_text((72, 100), "Battery Life | Battery Size", fontsize=11)
    page.insert_text((72, 120), "270 Hrs      | 13", fontsize=11)
    doc.save(path)
    doc.close()


def _write_scanned_like_pdf(path: Path) -> None:
    """Image-heavy page with minimal native text — triggers selective OCR path."""
    doc = pymupdf.open()
    page = doc.new_page()
    rect = pymupdf.Rect(72, 72, 400, 400)
    page.insert_image(rect, stream=_tiny_png())
    page.insert_text((72, 420), "Brochure Product Overview", fontsize=12)
    doc.save(path)
    doc.close()


def _tiny_png() -> bytes:
    # Minimal valid 1x1 PNG
    import base64

    return base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )


def _write_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("Company Policy", level=1)
    doc.add_paragraph("Return window: 10 days from delivery date.")
    doc.add_paragraph("All products must be returned in original packaging.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Policy"
    table.cell(0, 1).text = "Duration"
    table.cell(1, 0).text = "Returns"
    table.cell(1, 1).text = "10 days"
    doc.save(path)


def _write_product_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("Product Overview", level=1)
    doc.add_paragraph("Radius M16 BTE is a behind-the-ear hearing aid.")
    doc.add_heading("Technical Specifications", level=2)
    doc.add_paragraph("Battery Life: 270 Hrs")
    doc.add_paragraph("Battery Size: 13")
    doc.save(path)


def ensure_fixtures(base_dir: Path) -> None:
    base_dir.mkdir(parents=True, exist_ok=True)
    manifest = {
        "documents": [
            {"id": "digital_pdf", "filename": "digital_product.pdf", "document_type": "product"},
            {"id": "scanned_pdf", "filename": "scanned_brochure.pdf", "document_type": "brochure"},
            {"id": "brochure_pdf", "filename": "brochure_overview.pdf", "document_type": "brochure"},
            {"id": "docx_policy", "filename": "policy_returns.docx", "document_type": "policy"},
            {"id": "table_pdf", "filename": "product_specs_table.pdf", "document_type": "product"},
            {"id": "design_pdf", "filename": "design_brochure.pdf", "document_type": "brochure"},
            {"id": "product_docx", "filename": "radius_m16.docx", "document_type": "product"},
        ]
    }

    _write_pdf(
        base_dir / "digital_product.pdf",
        [
            "Bluup+ Product Guide\n\nBluup+ earplugs offer up to 29 dB noise reduction.",
            "Features\n\n- Advanced noise protection\n- Comfortable fit",
        ],
    )
    _write_scanned_like_pdf(base_dir / "scanned_brochure.pdf")
    _write_pdf(
        base_dir / "brochure_overview.pdf",
        ["Investor Brochure\n\nEarkart redefines hearing care across India."],
    )
    _write_docx(base_dir / "policy_returns.docx")
    _write_table_pdf(base_dir / "product_specs_table.pdf")
    _write_scanned_like_pdf(base_dir / "design_brochure.pdf")
    _write_product_docx(base_dir / "radius_m16.docx")

    (base_dir / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
