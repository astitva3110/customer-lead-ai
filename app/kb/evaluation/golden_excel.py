"""Load golden Q&A cases from Excel and resolve expected chunk IDs."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.helpers.text_match import grounding_fact_matches, normalize_match_text, split_answer_sentences
from app.kb.chunking.models import ProductionChunkRecord

GOLDEN_DATASET_VERSION = "golden_excel_v1"
PROJECT_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_GOLDEN_EXCEL_PATH = PROJECT_ROOT / "earkart_golden_dataset.xlsx"
DEFAULT_GOLDEN_SHEET = "Golden Dataset"
_PLACEHOLDER_RE = re.compile(r"\[FILL:", re.IGNORECASE)

_COLUMN_ALIASES: dict[str, set[str]] = {
    "id": {"id", "case_id", "qid", "question_id", "evaluation_id", "sr", "s.no", "s no", "no", "test id"},
    "question": {
        "question",
        "query",
        "user question",
        "user query",
        "user_question",
        "prompt",
        "questions",
    },
    "paraphrase": {"query variant / paraphrase", "query variant", "paraphrase"},
    "answer": {
        "answer",
        "expected_answer",
        "golden_answer",
        "ground_truth",
        "expected answer",
        "golden answer",
        "ideal answer (golden)",
        "ideal answer",
        "answers",
    },
    "grounding_points": {
        "must-have facts (grounding points)",
        "must-have facts",
        "grounding points",
        "must have facts",
    },
    "expected_source": {
        "expected source / kb doc",
        "expected source",
        "kb doc",
        "source",
    },
    "chunk_id": {"chunk_id", "expected_chunk_id", "expected_chunk_ids", "chunk ids"},
    "document_id": {"document_id", "doc_id", "expected_document_id", "document id"},
    "category": {"category", "topic", "type", "section"},
}


def _normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[_\-/]+", " ", text)
    text = re.sub(r"[()]+", " ", text)
    return " ".join(text.split())


def _normalized_alias_map() -> dict[str, str]:
    mapping: dict[str, str] = {}
    for field, aliases in _COLUMN_ALIASES.items():
        for alias in aliases:
            mapping[_normalize_header(alias)] = field
    return mapping


_NORMALIZED_ALIASES = _normalized_alias_map()


def _map_headers(headers: list[Any]) -> dict[str, int]:
    mapped: dict[str, int] = {}
    for index, header in enumerate(headers):
        normalized = _normalize_header(header)
        if not normalized:
            continue
        field = _NORMALIZED_ALIASES.get(normalized)
        if field and field not in mapped:
            mapped[field] = index
    return mapped


def _cell_text(row: tuple[Any, ...], index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    value = row[index]
    if value is None:
        return ""
    return str(value).strip()


def _parse_chunk_ids(raw: str) -> list[str]:
    if not raw:
        return []
    parts = re.split(r"[,;\n]+", raw)
    return [part.strip() for part in parts if part.strip()]


def _parse_grounding_points(raw: str) -> list[str]:
    if not raw:
        return []
    parts = re.split(r"[;\n]+", raw)
    needles: list[str] = []
    for part in parts:
        needle = part.strip()
        if not needle or _PLACEHOLDER_RE.search(needle):
            continue
        needles.append(needle)
    return needles


def _is_placeholder_text(text: str) -> bool:
    return bool(text and _PLACEHOLDER_RE.search(text))


def _resolve_sheet(workbook, sheet_name: str | None):
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Worksheet '{sheet_name}' not found. Available sheets: {workbook.sheetnames}"
            )
        return workbook[sheet_name]

    for candidate in (DEFAULT_GOLDEN_SHEET, "Golden Dataset", "Eval Runs"):
        if candidate in workbook.sheetnames:
            return workbook[candidate]
    return workbook.active


def load_golden_cases_from_excel(
    excel_path: Path,
    *,
    sheet_name: str | None = None,
) -> list[dict[str, Any]]:
    if not excel_path.exists():
        raise FileNotFoundError(
            f"Golden dataset Excel not found: {excel_path}. "
            f"Expected default: {DEFAULT_GOLDEN_EXCEL_PATH}"
        )

    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    worksheet = _resolve_sheet(workbook, sheet_name)
    rows = worksheet.iter_rows(values_only=True)
    header_row = next(rows, None)
    if not header_row:
        raise ValueError(f"Golden dataset Excel is empty: {excel_path}")

    column_map = _map_headers(list(header_row))
    if "question" not in column_map:
        raise ValueError(
            "Golden dataset Excel must include a question column "
            f"(found headers: {list(header_row)})."
        )

    cases: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows, start=2):
        question = _cell_text(row, column_map.get("question"))
        paraphrase = _cell_text(row, column_map.get("paraphrase"))
        answer = _cell_text(row, column_map.get("answer"))
        grounding_raw = _cell_text(row, column_map.get("grounding_points"))
        if not question and not answer and not grounding_raw:
            continue
        if not question:
            continue

        case_id = _cell_text(row, column_map.get("id")) or f"G{len(cases) + 1:03d}"
        chunk_ids = _parse_chunk_ids(_cell_text(row, column_map.get("chunk_id")))
        document_id = _cell_text(row, column_map.get("document_id")) or None
        category = _cell_text(row, column_map.get("category")) or "general"
        expected_source = _cell_text(row, column_map.get("expected_source")) or None
        grounding_points = _parse_grounding_points(grounding_raw)
        skipped_reason = None
        if _is_placeholder_text(answer) and not grounding_points and not chunk_ids:
            skipped_reason = "placeholder_answer"
        elif not answer and not grounding_points and not chunk_ids:
            skipped_reason = "missing_answer_and_grounding"

        cases.append(
            {
                "id": case_id,
                "query": question,
                "query_paraphrase": paraphrase or None,
                "expected_answer": answer,
                "grounding_points": grounding_points,
                "expected_source": expected_source,
                "expected_document_id": document_id,
                "expected_chunk_ids_override": chunk_ids,
                "category": category,
                "source_row": row_number,
                "source_sheet": worksheet.title,
                "skipped_reason": skipped_reason,
            }
        )

    workbook.close()
    if not cases:
        raise ValueError(f"No golden Q&A rows found in {excel_path}")
    return cases


def _resolve_from_needles(
    scope: list[ProductionChunkRecord],
    needles: list[str],
) -> tuple[list[str], bool]:
    if not needles:
        return [], False

    scored: list[tuple[int, str]] = []
    for chunk in scope:
        content = normalize_match_text(chunk.content)
        score = sum(1 for needle in needles if grounding_fact_matches(needle, content))
        if score > 0:
            scored.append((score, chunk.chunk_id))

    if not scored:
        return [], False

    scored.sort(key=lambda item: item[0], reverse=True)
    max_score = scored[0][0]
    best_ids = list(dict.fromkeys(chunk_id for score, chunk_id in scored if score == max_score))

    # Standard recall@k: pass if any expected chunk appears in top-k.
    if max_score == len(needles):
        single_chunk_matches = [
            chunk_id
            for chunk_id in best_ids
            if _chunk_matches_all_facts(
                next(chunk for chunk in scope if chunk.chunk_id == chunk_id),
                needles,
            )
        ]
        if single_chunk_matches:
            return single_chunk_matches[:1], False

    return best_ids, False


def _chunk_matches_all_facts(chunk: ProductionChunkRecord, needles: list[str]) -> bool:
    content = normalize_match_text(chunk.content)
    return all(grounding_fact_matches(needle, content) for needle in needles)


def resolve_expected_chunk_ids_from_answer(
    chunks: list[ProductionChunkRecord],
    *,
    answer: str,
    grounding_points: list[str] | None = None,
    document_id: str | None = None,
    chunk_ids_override: list[str] | None = None,
) -> tuple[list[str], bool]:
    """Resolve ground-truth chunk IDs from golden answer text and grounding points."""
    if chunk_ids_override:
        known = {chunk.chunk_id for chunk in chunks}
        resolved = [chunk_id for chunk_id in chunk_ids_override if chunk_id in known]
        return resolved, len(resolved) > 1

    scope = chunks
    if document_id:
        scope = [chunk for chunk in chunks if chunk.document_id == document_id]

    needle_ids, require_all = _resolve_from_needles(scope, grounding_points or [])
    if needle_ids:
        return needle_ids, require_all

    if _is_placeholder_text(answer):
        return [], False

    normalized_answer = normalize_match_text(answer)
    if not normalized_answer:
        return [], False

    matched = [
        chunk.chunk_id
        for chunk in scope
        if normalized_answer in normalize_match_text(chunk.content)
    ]
    if matched:
        unique_ids = list(dict.fromkeys(matched))
        return unique_ids, len(unique_ids) > 1

    if len(normalized_answer) > 80:
        anchor = normalized_answer[:80]
        anchor_matches = [
            chunk.chunk_id
            for chunk in scope
            if anchor in normalize_match_text(chunk.content)
        ]
        if anchor_matches:
            return [anchor_matches[0]], False

    per_sentence: list[str] = []
    for sentence in split_answer_sentences(answer):
        normalized_sentence = normalize_match_text(sentence)
        for chunk in scope:
            if normalized_sentence in normalize_match_text(chunk.content):
                per_sentence.append(chunk.chunk_id)
                break

    unique_sentence_ids = list(dict.fromkeys(per_sentence))
    if len(unique_sentence_ids) >= 2:
        return unique_sentence_ids, True
    if len(unique_sentence_ids) == 1:
        return unique_sentence_ids, False

    return [], False


def enrich_golden_cases(
    cases: list[dict[str, Any]],
    chunks: list[ProductionChunkRecord],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    evaluable: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []

    for case in cases:
        if case.get("skipped_reason"):
            skipped.append({**case, "expected_chunk_ids": [], "require_all_expected_chunks": False})
            continue

        expected_chunk_ids, require_all = resolve_expected_chunk_ids_from_answer(
            chunks,
            answer=case.get("expected_answer", ""),
            grounding_points=case.get("grounding_points") or None,
            document_id=case.get("expected_document_id"),
            chunk_ids_override=case.get("expected_chunk_ids_override") or None,
        )
        enriched = {
            **case,
            "expected_chunk_ids": expected_chunk_ids,
            "require_all_expected_chunks": require_all,
        }
        if expected_chunk_ids:
            evaluable.append(enriched)
        else:
            skipped.append({**enriched, "skipped_reason": "unresolved_chunk_ids"})

    return evaluable, skipped


def validate_golden_cases(cases: list[dict[str, Any]]) -> None:
    missing = [case["id"] for case in cases if not case.get("expected_chunk_ids")]
    if missing:
        raise ValueError(
            "Golden dataset has cases without resolvable expected_chunk_ids: "
            f"{missing[:10]}"
        )
